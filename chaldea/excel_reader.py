"""Đọc file Excel thành 2 dạng:

1. Văn bản dạng lưới (để LLM "nhìn" bảng): mỗi dòng ghi rõ số dòng Excel và tên cột,
   ô gộp được đánh dấu phạm vi, ví dụ  D5[D5:F5]="Some college or associate degree".
2. DataFrame cho từng sheet (để code Python tính toán): index = số dòng Excel,
   cột = chữ cái cột Excel, ô gộp được điền lại giá trị của ô góc trên-trái.

Giữ nguyên toạ độ Excel giúp LLM xử lý được tiêu đề nhiều tầng, nhiều bảng trong
một sheet mà không cần đoán cấu trúc.
"""
from __future__ import annotations

import datetime as dt
from dataclasses import dataclass
from functools import lru_cache

import openpyxl
import pandas as pd
from openpyxl.utils import get_column_letter


# Đo thật với Qwen3.7 Flash: bảng toàn số ~1.15-1.4 ký tự/token. 20K ký tự ≈ tối đa ~18K token/lần gọi,
# giữ chi phí mỗi câu dưới ~0.001 USD (giới hạn 0.10 USD/100 câu). File mẫu lớn nhất 17 978 ký tự -> không bị cắt.
MAX_TEXT_CHARS = 20_000
HEAD_SHARE = 0.75         # khi phải cắt: 75% dành cho phần đầu (tiêu đề + dữ liệu đầu), 25% cho phần cuối


@dataclass
class Workbook:
    text: str                          # bản mô tả gửi cho LLM
    sheets: dict[str, pd.DataFrame]    # dữ liệu để chạy code (luôn đầy đủ, không bị cắt)


def _used_extent(ws) -> tuple[int, int]:
    """Dòng/cột cuối thực sự có dữ liệu. ws.max_row có thể lên tới 1.048.576 nếu file
    bị định dạng cả cột -> đọc hết sẽ rất chậm và tốn RAM (giới hạn 1GB)."""
    try:
        used = [key for key, cell in ws._cells.items() if cell.value is not None]
    except AttributeError:
        return ws.max_row, ws.max_column
    if not used:
        return 0, 0
    return max(r for r, _ in used), max(c for _, c in used)


def _shorten(lines: list[str], budget: int) -> list[str]:
    """Cắt bớt phần giữa của một sheet quá dài: giữ phần đầu (tiêu đề, cấu trúc) và phần cuối.
    Code Python vẫn chạy trên DataFrame đầy đủ nên đáp án không bị ảnh hưởng."""
    if sum(len(x) + 1 for x in lines) <= budget:
        return lines
    head, used = [], 0
    for line in lines:
        if used + len(line) + 1 > budget * HEAD_SHARE:
            break
        head.append(line)
        used += len(line) + 1
    tail, used = [], 0
    for line in reversed(lines[len(head):]):
        if used + len(line) + 1 > budget * (1 - HEAD_SHARE):
            break
        tail.insert(0, line)
        used += len(line) + 1
    omitted = len(lines) - len(head) - len(tail)
    note = (f"... [{omitted} rows omitted here to save space; they follow the same layout. "
            f"ALWAYS compute over the full DataFrame in code, never only over the rows shown] ...")
    return head + [note] + tail


def _fmt(value) -> str:
    """Hiển thị một giá trị ô gọn gàng cho LLM."""
    if isinstance(value, dt.datetime):
        if value.time() == dt.time(0, 0):
            return value.date().isoformat()
        return value.isoformat(sep=" ")
    if isinstance(value, (dt.date, dt.time)):
        return value.isoformat()
    if isinstance(value, float):
        return repr(round(value, 10)).removesuffix(".0") if value.is_integer() else repr(round(value, 10))
    if isinstance(value, str):
        text = " ".join(value.split())  # gộp xuống dòng / khoảng trắng thừa
        return '"' + text.replace('"', "'") + '"'
    return str(value)


@lru_cache(maxsize=64)
def load_workbook(path: str) -> Workbook:
    """Đọc một file .xlsx (có cache vì nhiều câu hỏi dùng chung một file)."""
    wb = openpyxl.load_workbook(path, data_only=True)
    sheet_lines: list[list[str]] = []
    sheets: dict[str, pd.DataFrame] = {}

    for ws in wb.worksheets:
        max_row, max_col = _used_extent(ws)
        grid = [[ws.cell(r, c).value for c in range(1, max_col + 1)] for r in range(1, max_row + 1)]

        # Ô gộp: ghi nhãn phạm vi cho LLM + điền giá trị vào toàn vùng cho DataFrame
        merge_label: dict[tuple[int, int], str] = {}
        filled = [row[:] for row in grid]
        for rng in ws.merged_cells.ranges:
            top_val = grid[rng.min_row - 1][rng.min_col - 1]
            merge_label[(rng.min_row, rng.min_col)] = str(rng)
            for r in range(rng.min_row, rng.max_row + 1):
                for c in range(rng.min_col, rng.max_col + 1):
                    filled[r - 1][c - 1] = top_val

        cols = [get_column_letter(c) for c in range(1, max_col + 1)]
        df = pd.DataFrame(filled, index=range(1, max_row + 1), columns=cols)
        sheets[ws.title] = df

        lines = [f'=== Sheet "{ws.title}" (rows 1-{max_row}, columns A-{cols[-1] if cols else "A"}) ===']
        for r in range(1, max_row + 1):
            cells = []
            for c in range(1, max_col + 1):
                value = grid[r - 1][c - 1]
                if value is None or (isinstance(value, str) and not value.strip()):
                    continue
                ref = f"{cols[c - 1]}{r}"
                if (r, c) in merge_label:
                    ref += f"[{merge_label[(r, c)]}]"
                cells.append(f"{ref}={_fmt(value)}")
            if cells:
                lines.append(" | ".join(cells))
        sheet_lines.append(lines)

    # Tổng quá dài -> chia ngân sách: sheet nhỏ giữ nguyên, phần còn lại chia đều cho các sheet lớn
    sizes = [sum(len(x) + 1 for x in lines) for lines in sheet_lines]
    if sum(sizes) > MAX_TEXT_CHARS:
        budget, remaining = {}, MAX_TEXT_CHARS
        order = sorted(range(len(sizes)), key=sizes.__getitem__)
        for n, i in enumerate(order):
            budget[i] = min(sizes[i], remaining // (len(order) - n))
            remaining -= budget[i]
        sheet_lines = [_shorten(lines, budget[i]) for i, lines in enumerate(sheet_lines)]
    return Workbook(text="\n\n".join("\n".join(lines) for lines in sheet_lines), sheets=sheets)
